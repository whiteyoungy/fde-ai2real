#!/usr/bin/env python3
"""从五份课件里抽出每一条例子，生成对应的微课写作提示词。

为什么要从 pptx 里抽而不是从 dayN.py 里抽：
    dayN.py 里的例子是散在各处的函数调用，抽出来拿不到它在页面上的邻居
    （标题、表格、正文），而那些邻居正是提示词最需要的上下文——
    没有它们，模型不知道这条例子是在论证什么。pptx 里所有内容都带坐标，
    按 y 排一遍就还原出了这一页的阅读顺序。

分档规则见 tier()，阈值是照着实际分布调出来的，不是拍的。
"""
import glob
import pathlib
import re
from pptx import Presentation

HERE = pathlib.Path(__file__).resolve().parent
SLIDES = HERE.parent / "slides"
EMU = 914400.0

DAYS = {
    "第1天-风险关卡与前两关": (1, "风险、关卡、死法、选题、定界", "高管 ＋ FDE"),
    "第2天-验收采用合规与现场": (2, "验收、采用、合规、计价，以及从关卡走到现场", "上午高管 ＋ FDE，下午 FDE"),
    "第3天-工程基线到本体落地": (3, "接得上、数得清：从工程基线到本体落地", "FDE"),
    "第4天-护栏评估与国产化": (4, "让它可信：护栏、评估与国产化", "FDE"),
    "第5天-读一个开源项目": (5, "读一个开源项目，判断它能不能进你的交付", "FDE"),
}

# 模块 → （全称，挂在骨架哪一层）
MODULES = {
    "00": ("开场与收束", "全局"),
    "M1": ("这个岗位在解决什么问题", "公理 · 风险转移"),
    "M2": ("六道关卡", "过程 · 六道关卡"),
    "M3": ("怎么把自己做死", "结果 · 接管（四个陷阱都在毁掉它）"),
    "M4": ("第 1 关 选题", "过程 · 第 1 关"),
    "M5": ("第 2 关 定界", "过程 · 第 2 关"),
    "M6": ("第 4 关 验收", "结果 · 证据"),
    "M7": ("第 6 关 采用", "结果 · 接管"),
    "M8": ("合规与备案", "过程 · 第 2 关的硬约束"),
    "M9": ("计价方式就是责任分配", "公理 · 风险转移在商务条款上的体现"),
    "M10": ("第 3 关 切片", "过程 · 第 3 关"),
    "M11": ("最小可执行业务模型", "资产 · 七要素"),
    "M12": ("第 5 关 上线", "结果 · 控制"),
    "M13": ("国企、银行、政务与驻场", "过程 · 现场制度约束"),
    "M14": ("经验资产化", "资产 · 底座与现场层"),
    "M15": ("生产级工程基线", "结果 · 证据（代码本身要能被检验）"),
    "M16": ("规约驱动开发与工具选型", "结果 · 证据（把“什么算对”搬进仓库）"),
    "M17": ("企业集成与协议桥接", "结果 · 证据（接不上就没有真实运行记录）"),
    "M18": ("数据管道与向量层", "结果 · 证据（语料质量决定记录有没有意义）"),
    "M19": ("本体的工程实现", "资产 · 七要素落层"),
    "M20": ("生成式系统与智能体编排", "过程 · 第 3 关到第 5 关"),
    "M21": ("部署、护栏与可观测性", "结果 · 控制"),
    "M22": ("评估的工程实现", "结果 · 证据 ＋ 资产 · 验收集"),
    "M23": ("国产模型与信创改造", "过程 · 第 2 关的硬约束"),
    "M24": ("四个案例串讲", "全部四层"),
    "M25": ("90 天上岗与作品集", "全局"),
    "M26": ("一棵可以拆开的树", "过程 · 交付形态"),
    "M27": ("接得上，才谈得上别的", "过程 · 第 3 关的前置条件"),
    "M28": ("护栏会自报强制力", "结果 · 控制"),
    "M29": ("模型可见即已记录", "结果 · 证据"),
    "M30": ("从会话日志里长出验收集", "资产 · 验收集"),
    "M31": ("把它交出去 · 什么时候不该用它", "公理 · 风险转移"),
    "M32": ("另一个样本：Semantica", "资产 · 七要素覆盖度"),
    "上机 ①": ("上机 · Lab-11 工作流发现", "过程 · 第 1、2 关"),
    "上机 ②": ("上机 · Lab-04 混合检索", "结果 · 证据"),
    "上机 ③": ("上机 · Lab-05 门禁", "结果 · 证据 ＋ 资产"),
    "上机 ④": ("上机 · 开源组件进场评估", "全部四层"),
    "研讨": ("决策复盘研讨", "全部四层"),
    "收尾": ("高管专场收尾", "全局"),
    "M13 · M14": ("驻场与资产化", "过程 · 现场 ＋ 资产"),
    "M6 · M7": ("验收与采用", "结果 · 证据 ＋ 接管"),
}

# ── 分档信号 ──────────────────────────────────────────────
SRC_HARD = re.compile(
    r"verify\.sh|实测|官方|论文|年报|审计报告|SEC|10-K|逐字|核对|条文|原文|"
    r"招聘启事|模型卡|研究|基准报告|判决|裁决")
NUM = re.compile(
    r"\d+\s*%|\d+\s*/\s*\d+|\d[\d,\.]*\s*(万|亿|条|页|倍|人月|美元|毫秒|"
    r"分钟|小时|天|周|个月|年|字节|GB|MB|tok)")
XREF_ONLY = re.compile(r"^(这条|这和|这就是|这正是|这一条|它和|同一条|同样|"
                       r"这也是|这张表|这四条|这两条|回到|对照)")
# 有叙事：有具体的人、机构、时间点——能撑起一段完整的故事
NARR = re.compile(
    r"某[一二三四五六七八九十百千]?[家位条个]?[银企公市医大产客工]|一家|一位|"
    r"当时|事后|上线[后前]|复盘|那次|那一|第 ?\d+ ?[周天月年]|年[初底]|三个月")
# 纯导航：预告别处会讲、指回讲义——自己不带新信息
NAV = re.compile(
    r"第 ?[1-5一二三四五] ?天|明天|昨天|回去|会展开|会讲|那一节|这一节|"
    r"下一页|上一页|讲义里|见教材|完整版")


# 课程结构引用：“第 2 天”“第 3 关”“§9.5”这类是导航，不是数据。
# 打分前先剔掉，否则一句纯呼应会因为里面有个“第 2 天”被算成有数字。
STRUCT = re.compile(r"第 ?[1-5一二三四五] ?天|第 ?\d+ ?[章关页层步条次种类项组行部]|"
                    r"§ ?\d+(\.\d+)?|上册|下册|别册[一二]|Lab-\d+|附录 ?[A-E]")


def tier(body, source):
    """三档：重 / 中 / 轻。

    阈值不是拍的，是照实际分布调的：重 >= 2 得 66 条，中 >= 0 得 198 条，
    其余 152 条落到轻档。分数含义大致是——
    有可查证来源加 2，有数字加 2，有叙事加 1，够长加 1；
    纯呼应扣 2，纯导航扣 1，太短扣 1。

    第 5 天只有 1 条落进重档，这不是分档器的 bug，是那一天的写法使然：
    第 5 天的例子大多是把当天的判断挂回前四天讲过的东西，
    真正撑得起一节完整微课的材料在页面主体（对照表、判据卡）里，不在例子卡上。
    """
    plain = STRUCT.sub("", body)      # 剔掉课程结构引用再看有没有真数字
    score = 0
    if SRC_HARD.search(source):
        score += 2
    if NUM.search(plain):
        score += 2
    if SRC_HARD.search(body):
        score += 1
    if len(body) >= 110:
        score += 1
    elif len(body) < 70:
        score -= 1
    if NARR.search(plain):
        score += 1
    if len(NUM.findall(plain)) >= 2:
        score += 1
    # 纯呼应型：只是把这一条挂回前面讲过的东西，自己不带新信息
    if XREF_ONLY.match(body) and not NUM.search(plain) and len(body) < 100:
        score -= 2
    # 纯导航型：说的是别处会讲什么
    if NAV.search(body) and not NUM.search(plain):
        score -= 1
    return "重" if score >= 2 else ("中" if score >= 0 else "轻")


TIER_SPEC = {
    "重": ("6–8 分钟", "900–1200 字"),
    "中": ("3–4 分钟", "450–650 字"),
    "轻": ("1–2 分钟，或并入邻近微课", "180–260 字"),
}


# ── 抽取 ──────────────────────────────────────────────────

def blocks(slide):
    """按 y 排序还原这一页的阅读顺序。"""
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if not t:
            continue
        out.append((sh.top / EMU, sh.left / EMU, sh.width / EMU, t))
    out.sort(key=lambda b: (round(b[0], 2), b[1]))
    return out


def parse_slide(slide):
    bs = blocks(slide)
    eyebrow = heading = footer = ""
    examples = []       # (y, body, source)
    others = []         # (y, text)
    for y, x, w, t in bs:
        if y < 0.8 and x < 2.0:
            eyebrow = t
        elif y < 0.8:
            continue                       # 右上角的“第 N 天”标
        elif y > 6.7:
            footer = t
        elif not heading and y < 1.9 and x < 1.3:
            heading = t
        elif t.startswith("例　"):
            parts = t.split("\n")
            src = ""
            if len(parts) > 1 and parts[-1].startswith("— "):
                src = parts[-1][2:].strip()
                parts = parts[:-1]
            examples.append((y, "\n".join(parts)[2:].strip(), src))
        else:
            others.append((y, t))
    return eyebrow, heading, footer, examples, others


def module_of(eyebrow):
    """页眉形如 'M1   1.3 问责' 或 '上机 ①　　15:00 · 90 分钟'。"""
    e = eyebrow.replace("　", " ")
    for key in sorted(MODULES, key=len, reverse=True):
        if e.startswith(key):
            label = e[len(key):].strip()
            return key, label
    return "", e


CTX_LIMIT = {"重": 560, "中": 360, "轻": 200}


def excerpt(body, n=28):
    """取例子开头一段当标题。

    不能简单截断——课件正文里全角引号成对出现，截在引号中间，
    索引读起来就是半句话。所以截完要把没配对的那个引号处理掉。
    """
    s = body.split("。")[0].strip()
    if len(s) < 10 and "。" in body:
        s = "。".join(body.split("。")[:2]).strip()
    if len(s) > n:
        s = s[:n]
        cut = max(s.rfind(c) for c in "，；：、——")
        if cut >= n - 10:
            s = s[:cut]
        s = s.rstrip("，；：、—") + "…"
    if s.count("“") > s.count("”"):
        s = s[:s.rfind("“")].rstrip("，；：、—")
        if not s.endswith("…"):
            s += "…"
    return s


def compress(others, limit=520):
    """把这一页除例子外的内容压成一段，给模型当上下文。

    按档位截不同长度：轻档只要 180–260 字的补充口播，
    塞一整张时间表进去只会挤掉真正的信息。
    """
    txt = "；".join(t.replace("\n", " / ") for _, t in others)
    txt = re.sub(r"\s+", " ", txt)
    return txt[:limit] + ("…" if len(txt) > limit else "")


# ── 生成 ──────────────────────────────────────────────────

PROMPT = """# {code} · {title}

<!-- 档位：{tierc}　|　第 {day} 天 {mod} {label}　|　课件第 {page} 页　|　\
目标 {dur} -->

你是一位为企业内训写微课脚本的作者。这门课是《前线部署工程师（FDE）：把 AI 交付到真实世界》
的五天面授版，学员是{aud}。请按下面的材料和要求，写一份微课脚本。

---

## 材料

### 一、这条微课要讲的例子

{body}

**来源**：{source}

### 二、它在课上的位置

| | |
|---|---|
| 第几天 | 第 {day} 天：{daytitle} |
| 所在模块 | {mod} {modname} |
| 所在小节 | {label} |
| 挂在骨架哪一层 | {layer} |
| 课件页码 | 第 {page} 页 |

### 三、这一页在论证什么

**这一页的主张**：{heading}

**这一页的其他内容**：{context}
{footline}
### 四、前后文

- **前一页**：{prev}
- **后一页**：{nxt}

---

## 要求

**时长**：{dur}　**正文字数**：{words}

{tierline}

写作时守住四条：

1. **从例子进，不从定义进。** 开头三十秒之内要出现这个例子的具体场景，
   不要先讲一段概念再举例。学员点开这条微课，是因为它的标题写着一个具体的事。
2. **把判据说清楚。** 微课的落点不是“记住这个故事”，是“下次遇到同类情况，
   用哪一句话做判断”。脚本结尾要能提炼出一句可以直接用的判据。
3. **来源怎么说就怎么说。** 上面“来源”那一栏里如果写了“实测”“官方”“原文”，
   脚本里可以说得肯定；如果写了“归纳”“推导”“教学示例”“未核实”“合成案例”，
   脚本里必须照实说明，不能升级成确定结论。
4. **不要新编案例。** 需要补充例证时，只能用这门课已有的材料
   （六道关卡、12 个 Lab 的实测数字、四个案例章、两个决策复盘）。

## 输出格式

```
【标题】     一句话，不超过 18 字，要具体不要笼统
【一句话结论】 这条微课要让学员带走的那一句判据
【时长】     预估分钟数

【口播脚本】
（分段写，每段前面标一个时间戳，比如 00:00 / 00:40 / 01:30。
  口语，短句，可以直接照着念。不要出现“接下来我们来看”这类过场话。）

【画面】
（逐段给画面说明：这一段屏幕上应该出现什么。
  可以是课件的哪一页、一张什么表、一个什么数字、一段什么对照。
  不需要设计动效，说清楚画面上有什么就行。）

【随堂测题】
（一道单选或判断题，考的是判据不是记忆。给出正确答案和一句为什么。）

【延伸】
（这条微课往下追可以看教材的哪一节，一行。）
```

## 红线

- 不引用任何竞品的案例或文字。
- 不新编数字。脚本里出现的每个数字都要能在上面的材料里找到。
- 中文正文用全角引号“”，不要用「」，不要用 ASCII 引号。
- 不写“众所周知”“显而易见”“毫无疑问”这类填充词。
"""

TIERLINE = {
    "重": "这一条是**重档**：例子本身带完整案例、可查证数字或官方原文，"
          "值得单独成课。脚本要把来龙去脉讲完整——事情怎么发生的、"
          "当时可以怎么判断、判错了代价是什么。",
    "中": "这一条是**中档**：例子有实质内容但不够撑满一节完整微课。"
          "脚本走精简结构——场景、判据、一句总结，不铺陈背景。",
    "轻": "这一条是**轻档**：例子本身是一句呼应或一条补充说明。"
          "**先判断它值不值得单独成课**——如果它离不开上下文，"
          "就在输出的开头写一行“建议并入：<同一模块里哪一条微课>”，"
          "然后只写一段 180–260 字的补充口播，供那条微课在结尾处使用。",
}


def main():
    stats = {"重": 0, "中": 0, "轻": 0}
    total = 0
    index = {}

    for path in sorted(SLIDES.glob("*.pptx")):
        stem = path.stem
        if stem not in DAYS:
            continue
        day, daytitle, aud = DAYS[stem]
        prs = Presentation(str(path))
        parsed = [parse_slide(s) for s in prs.slides]
        headings = [p[1] for p in parsed]

        outdir = HERE / ("day%d" % day)
        outdir.mkdir(exist_ok=True)
        for old in outdir.glob("d%d-*.md" % day):
            old.unlink()

        rows = []
        seq = 0
        for i, (eyebrow, heading, footer, examples, others) in enumerate(parsed):
            if not examples:
                continue
            mod, label = module_of(eyebrow)
            modname, layer = MODULES.get(mod, ("", "全局"))
            page = i + 1
            prev = headings[i - 1] if i > 0 and headings[i - 1] else "（分隔页）"
            nxt = headings[i + 1] if i + 1 < len(headings) and headings[i + 1] else "（分隔页）"
            for _, body, source in examples:
                seq += 1
                t = tier(body, source)
                stats[t] += 1
                total += 1
                code = "D%d-%03d" % (day, seq)
                # 标题用例子本身的开头，不用页面标题——
                # 索引是拿来找“哪条微课讲什么”的，页面标题在元数据里已经有了
                title = excerpt(body)
                dur, words = TIER_SPEC[t]
                fname = "%s-%s.md" % (code.lower(), t)
                (outdir / fname).write_text(PROMPT.format(
                    code=code, title=title, tierc=t, day=day, mod=mod,
                    label=label, page=page, dur=dur, aud=aud,
                    body=body, source=source or "（课件未标注来源，写作前请回教材核对）",
                    daytitle=daytitle, modname=modname, layer=layer,
                    heading=heading or "（本页无标题）",
                    context=compress(others, CTX_LIMIT[t]) or "（本页只有这条例子）",
                    footline=("\n**页脚**：%s\n" % footer) if footer else "",
                    prev=prev, nxt=nxt, words=words,
                    tierline=TIERLINE[t]), encoding="utf-8")
                rows.append((code, t, page, mod, label, title, fname,
                             (heading or "")[:26]))
        index[day] = (daytitle, aud, rows)

    for day, (daytitle, aud, rows) in index.items():
        lines = ["# 第 %d 天 微课索引" % day, "",
                 "**主题**：%s　**听众**：%s" % (daytitle, aud), "",
                 "共 %d 条。重 %d ／ 中 %d ／ 轻 %d。" % (
                     len(rows),
                     sum(1 for r in rows if r[1] == "重"),
                     sum(1 for r in rows if r[1] == "中"),
                     sum(1 for r in rows if r[1] == "轻")), "",
                 "| 编号 | 档 | 页 | 模块 | 这条讲什么 | 所在页的主张 |",
                 "|---|---|---|---|---|---|"]
        for code, t, page, mod, label, title, fname, hd in rows:
            lines.append("| [%s](%s) | %s | %d | %s %s | %s | %s |"
                         % (code, fname, t, page, mod, label, title, hd))
        (HERE / ("day%d" % day) / "INDEX.md").write_text(
            "\n".join(lines) + "\n", encoding="utf-8")

    write_readme(index, stats, total)
    print("共 %d 条：重 %d ／ 中 %d ／ 轻 %d"
          % (total, stats["重"], stats["中"], stats["轻"]))
    for day, (_, _, rows) in index.items():
        print("  第 %d 天 %3d 条（重 %2d 中 %2d 轻 %2d）" % (
            day, len(rows),
            sum(1 for r in rows if r[1] == "重"),
            sum(1 for r in rows if r[1] == "中"),
            sum(1 for r in rows if r[1] == "轻")))


README = """# 微课提示词

《前线部署工程师（FDE）：把 AI 交付到真实世界》五天面授课，课件里每一条例子对应一份
微课写作提示词。**这里放的是提示词，不是微课本身**——每份文件都是自足的，
整份粘给大模型就能出一份微课脚本。

## 怎么用

1. 打开 `dayN/INDEX.md`，按“这条讲什么”那一列找到要做的那条。
2. 把对应的 `dN-xxx-档.md` **整份**内容粘给模型。
3. 拿到脚本之后按【画面】那一段去配课件页、录屏或录音。

提示词是从课件 pptx 里自动抽的。课件改了就重跑：

```bash
python3 microlec/build_prompts.py
```

## 三档怎么分

| 档 | 条数 | 目标时长 | 正文字数 | 什么样的例子落进来 |
|---|---|---|---|---|
| **重** | {n_h} | 6–8 分钟 | 900–1200 字 | 带完整案例、可查证数字或官方原文 |
| **中** | {n_m} | 3–4 分钟 | 450–650 字 | 有实质内容，但撑不满一节完整微课 |
| **轻** | {n_l} | 1–2 分钟，或并入邻近微课 | 180–260 字 | 一句呼应或一条补充说明 |

分档是打分打出来的，不是人工挑的（规则见 `build_prompts.py` 的 `tier()`）：
有可查证来源加 2、有数字加 2、有叙事加 1、够长加 1；纯呼应扣 2、纯导航扣 1、
太短扣 1。打分前先剔掉“第 2 天”“§9.5”这类课程结构引用——
它们是导航，不是数据。

**轻档的提示词不强求单独成课。** 它先让模型判断这条值不值得独立，
不值得就输出一行“建议并入：<哪一条>”，再给一段可以接在那条微课结尾的补充口播。

## 各天分布

| 天 | 主题 | 总数 | 重 | 中 | 轻 |
|---|---|---|---|---|---|
{rows}

**第 5 天只有 1 条重档，这不是分档器的问题。** 那一天的例子大多是把当天的判断
挂回前四天讲过的东西——真正撑得起一节完整微课的材料在页面主体
（可配置 vs 可替换那张对照表、护栏自报强制力那张表、七要素覆盖度表）里，
不在例子卡上。要给第 5 天做重档微课，素材应该从课件页面取，不是从例子卡取。

## 每份提示词里有什么

- **例子原文与来源**——来源那一栏决定了脚本能把话说多满
- **它在课上的位置**——第几天、哪个模块、挂在四层骨架的哪一层、课件第几页
- **这一页在论证什么**——页面主张 ＋ 同页的表格和正文（按档位截不同长度）
- **前后文**——前一页和后一页各讲了什么
- **写作要求与输出格式**——口播脚本、逐段画面说明、一道随堂测题
- **红线**——不引竞品、不新编数字、中文用全角引号

## 四条写作纪律（每份提示词里都写着）

1. 从例子进，不从定义进。
2. 把判据说清楚——落点是“下次怎么判断”，不是“记住这个故事”。
3. 来源怎么说就怎么说。标着“归纳”“推导”“教学示例”的，脚本里不能升级成确定结论。
4. 不新编案例。补充例证只能用这门课已有的材料。

## 目录

```
microlec/
  README.md            本文件
  build_prompts.py     生成器：从 slides/*.pptx 抽例子、分档、出提示词
  day1/ … day5/
    INDEX.md           这一天的索引：编号、档位、页码、模块、这条讲什么
    d1-001-轻.md       一条一份，文件名带档位
    …
```
"""


def write_readme(index, stats, total):
    rows = []
    for day in sorted(index):
        daytitle, _, rs = index[day]
        rows.append("| [第 %d 天](day%d/INDEX.md) | %s | %d | %d | %d | %d |" % (
            day, day, daytitle, len(rs),
            sum(1 for r in rs if r[1] == "重"),
            sum(1 for r in rs if r[1] == "中"),
            sum(1 for r in rs if r[1] == "轻")))
    rows.append("| **合计** | | **%d** | **%d** | **%d** | **%d** |"
                % (total, stats["重"], stats["中"], stats["轻"]))
    (HERE / "README.md").write_text(README.format(
        n_h=stats["重"], n_m=stats["中"], n_l=stats["轻"],
        rows="\n".join(rows)), encoding="utf-8")


if __name__ == "__main__":
    main()
