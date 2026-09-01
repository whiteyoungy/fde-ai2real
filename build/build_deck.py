#!/usr/bin/env python3
# SPDX-License-Identifier: MulanPSL-2.0
# Copyright (c) 2026 whiteyoungy
"""生成《前线部署工程师（FDE）：把 AI 交付到真实世界》内容导览 pptx。

视觉沿用 HTML 版：冷调纸白底、墨色正文、青色表示已实测、琥珀表示存疑。
所有数字用等宽字体——这本书每条结论都带数字，让它们成为视觉主角。
"""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

GROUND = RGBColor(0xF3, 0xF5, 0xF2)
PANEL  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x14, 0x1C, 0x24)
INK2   = RGBColor(0x38, 0x49, 0x5A)
MUTED  = RGBColor(0x6B, 0x7E, 0x8F)
RULE   = RGBColor(0xD6, 0xDD, 0xD8)
VERIFY = RGBColor(0x0B, 0x8E, 0x77)
CAVEAT = RGBColor(0xB2, 0x6E, 0x10)

CN = "微软雅黑"          # 中文字体：Windows/WPS 通用
MONO = "Consolas"       # 数字与代码

prs = Presentation()
prs.slide_width, prs.slide_height = In(13.333), In(7.5)
W, H = 13.333, 7.5
ML = 0.95               # 左边距


def set_cjk(run, cn=CN):
    """python-pptx 只写 latin 字体，中文会走回退。手动补 eastAsia。"""
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if ea is None:
        ea = etree.SubElement(rPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    ea.set('typeface', cn)


def slide(bg=GROUND):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    # 左侧细轨，呼应 HTML 版的进度条
    bar = s.shapes.add_shape(1, In(0), In(0), In(0.055), In(H))
    bar.fill.solid(); bar.fill.fore_color.rgb = VERIFY
    bar.line.fill.background(); bar.shadow.inherit = False
    return s


def tb(s, x, y, w, h):
    box = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def line(tf, text, size, color=INK, bold=False, mono=False,
         space_after=0, space_before=0, first=False, spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = MONO if mono else CN
    set_cjk(r, MONO if mono else CN)
    return p


def eyebrow(s, num, label):
    tf = tb(s, ML, 0.62, 9, 0.35)
    p = tf.paragraphs[0]
    if num:
        r = p.add_run(); r.text = num + "   "
        r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = VERIFY
        r.font.name = MONO; set_cjk(r, MONO)
    r = p.add_run(); r.text = label
    r.font.size = Pt(12); r.font.color.rgb = MUTED
    r.font.name = CN; set_cjk(r)


def heading(s, text, y=1.15, size=34, w=11.3):
    tf = tb(s, ML, y, w, 1.5)
    line(tf, text, size, INK, bold=True, first=True, spacing=1.15)


def body(s, text, y, size=16, color=INK2, w=10.6, spacing=1.5):
    tf = tb(s, ML, y, w, 1.4)
    line(tf, text, size, color, first=True, spacing=spacing)
    return tf


def bignum(s, x, y, value, unit, color=VERIFY, nsize=54, w=5.4):
    tf = tb(s, x, y, w, 1.0)
    line(tf, value, nsize, color, bold=True, mono=True, first=True, spacing=1.0)
    if unit:
        line(tf, unit, 12.5, MUTED, space_before=4, spacing=1.3)


def card(s, x, y, w, h, kicker, title, text, accent=None):
    box = s.shapes.add_shape(1, In(x), In(y), In(w), In(h))
    box.fill.solid(); box.fill.fore_color.rgb = PANEL
    box.line.color.rgb = accent or RULE
    box.line.width = Pt(1); box.shadow.inherit = False
    tf = tb(s, x + 0.26, y + 0.22, w - 0.5, h - 0.4)
    first = True
    if kicker:
        line(tf, kicker, 10.5, accent or MUTED, bold=True, first=True, space_after=5)
        first = False
    if title:
        line(tf, title, 15, INK, bold=True, first=first, space_after=5, spacing=1.25)
        first = False
    if text:
        line(tf, text, 12.5, INK2, first=first, spacing=1.4)


def table(s, x, y, rows, widths, head=True, hl=None, fsize=13):
    """rows[0] 为表头。hl 为需要高亮的行索引。"""
    cy = y
    for i, row in enumerate(rows):
        cx = x
        is_head = head and i == 0
        for j, cell in enumerate(row):
            tf = tb(s, cx, cy, widths[j], 0.3)
            col = MUTED if is_head else (VERIFY if hl == i else INK2)
            mono = (not is_head) and any(ch.isdigit() for ch in cell) and j > 0
            line(tf, cell, 10.5 if is_head else fsize, col,
                 bold=is_head or hl == i, mono=mono, first=True)
            cx += widths[j]
        cy += 0.34 if is_head else 0.40
        ln = s.shapes.add_shape(1, In(x), In(cy - 0.08), In(sum(widths)), In(0.012))
        ln.fill.solid(); ln.fill.fore_color.rgb = RULE
        ln.line.fill.background(); ln.shadow.inherit = False
    return cy


def quote(s, y, text, w=10.2):
    bar = s.shapes.add_shape(1, In(ML), In(y), In(0.045), In(0.78))
    bar.fill.solid(); bar.fill.fore_color.rgb = VERIFY
    bar.line.fill.background(); bar.shadow.inherit = False
    tf = tb(s, ML + 0.28, y + 0.02, w, 0.8)
    line(tf, text, 14, INK2, first=True, spacing=1.45)


# ══════════════ 1 封面 ══════════════
s = slide(PANEL)
eyebrow(s, None, "FORWARD DEPLOYED ENGINEER")
tf = tb(s, ML, 1.55, 11, 2.4)
line(tf, "把 AI 真正交付到", 52, INK, bold=True, first=True, spacing=1.12)
line(tf, "客户现场的那套手艺", 52, INK, bold=True, spacing=1.12)
body(s, "上下两册 524 页、另有三本别册的中文技术书。它和同类书最大的区别是：书里的结论都量过——"
        "12 个动手 Lab 各带一份验收脚本，全项通过才允许写进正文。", 4.25, size=17)
ln = s.shapes.add_shape(1, In(ML), In(5.35), In(11.2), In(0.012))
ln.fill.solid(); ln.fill.fore_color.rgb = RULE
ln.line.fill.background(); ln.shadow.inherit = False
for i, (v, k) in enumerate([("32", "章正文"), ("12", "动手 Lab"),
                            ("524", "页 A4"), ("530", "条引文核对")]):
    tf = tb(s, ML + i * 2.5, 5.6, 2.3, 0.9)
    line(tf, v, 30, INK, bold=True, mono=True, first=True)
    line(tf, k, 11, MUTED, space_before=2)

# ══════════════ 2 角色定位 ══════════════
s = slide(); eyebrow(s, "01", "角色定位")
heading(s, "FDE 不是售前，不是实施，也不是驻场运维")
body(s, "这个角色的特殊之处在于：它同时要对「这东西能不能用」和「客户愿不愿意用」负责。"
        "技术做对了但没人用，和压根没做出来，在结果上没有区别。", 2.35)
for i, (k, t) in enumerate([
        ("和售前的区别", "售前交付的是承诺，FDE 交付的是能在客户环境里跑起来的东西。"),
        ("和实施的区别", "实施按既定方案落地，FDE 要先把问题定义出来——客户说的需求往往不是真需求。"),
        ("和运维的区别", "运维保系统活着，FDE 要让它被用起来，并且用出业务结果。")]):
    card(s, ML + i * 3.8, 3.7, 3.5, 1.9, k, None, t)

# ══════════════ 3 四层骨架 ══════════════
s = slide(); eyebrow(s, "02", "骨架")
heading(s, "全书挂在四层骨架上，每一层回答一个问题", size=32)
body(s, "六道关卡是用得最多的一层，但它只是骨架的第二层。四层各自对得上一支成熟的研究传统——"
        "书里逐层标了出处档位，哪几个名字是本书自己造的也写在明处。", 2.2, size=15)
table(s, ML, 3.05, [
    ["层", "回答什么", "一句话"],
    ["公理", "为什么有这个岗位", "风险转移——买的不是代码，是把「能不能在我的场景落地」这个不确定性转移给一个担责的人"],
    ["过程", "按什么顺序走", "六道关卡：选题 → 定界 → 切片 → 验收 → 上线 → 采用"],
    ["结果", "走到哪儿算到", "证据给业务、控制给安全、接管给自己。三样各自对着一个会说「不行」的人"],
    ["资产", "什么留得下来", "验收集——模型、供应商、芯片都换过之后，只有它不用重建"],
], [1.1, 2.3, 7.6], fsize=13)
quote(s, 6.3, "全书的命题比岗位大一圈：AI 没有制造新的组织问题，它把老问题从「可以绕过」推到了「绕不过」——"
               "过去写在制度里的管理规则，现在必须被编译进系统本身。六关是这条命题在交付上的形状。")

# ══════════════ 4 六道关卡 ══════════════
s = slide(); eyebrow(s, "03", "方法论")
heading(s, "六道关卡，每一关的通过标准都能用「是 / 否」回答", size=30)
body(s, "关键不在于分了几关，而在于每一关的通过标准必须可判定——"
        "「需求明确了」不能作为通过标准，因为没人能回答明确到什么程度算明确。", 2.2, size=15)
table(s, ML, 3.05, [
    ["关卡", "通过标准（节选）"],
    ["选题", "至少 3 场情境访谈，每场只跟 1 位一线操作者；问题陈述句经受访者当面复述确认"],
    ["定界", "客户签字的范围文档，含 ≤3 条量化验收指标，假设条款须客户方书面确认"],
    ["切片", "最窄路径用真实数据端到端跑通至少一次，客户方能现场复现；技术债登记并设偿还期限"],
    ["验收", "评估集含 ≥20 条真实历史用例（含边缘与失败案例），分版本指标经客户技术负责人签字"],
    ["上线", "生产就绪清单全部通过（含人工复核队列、静默失败监控、值班表）；小流量观察窗口无严重事故"],
    ["采用", "连续 ≥4 周采用率达到约定阈值；一次跨部门种子用户复盘会；客户团队独立完成过一次常规变更并自己验证"],
], [1.5, 9.5], fsize=13)
quote(s, 6.35, "「客户签字的范围文档」能回答（有签字 / 没签字），"
                "「客户满意了」回答不了。这是判断通过标准合不合格最快的方法。")

# ══════════════ 4 混合检索 ══════════════
s = slide(); eyebrow(s, "04", "实测结论 · 检索")
heading(s, "「混合检索比纯向量好」——先别信，这句话有前提")
bignum(s, ML, 2.25, "98%", None, VERIFY, 50)
tf = tb(s, ML + 2.4, 2.42, 8.4, 1.0)
line(tf, "RRF 混合的 Recall@5。但同一份语料上，换一种构造方式，"
         "纯向量能到 100%，混合反而更差。", 15, INK2, first=True, spacing=1.4)
table(s, ML, 3.65, [
    ["方案", "总分", "语义改写类", "埋编号类"],
    ["BM25 单路", "92%", "3/8", "40/40"],
    ["纯向量", "90%", "8/8", "34/40"],
    ["RRF 混合", "98%", "7/8", "40/40"],
], [3.0, 1.6, 2.0, 1.8], hl=3)
body(s, "混合能赢的唯一前提是两个检索器的失败集不相交。注意「语义改写类」那一列——"
        "混合 7/8，比纯向量的 8/8 还低。融合不是免费的，总账赢了才叫赢。", 5.85, size=14.5)

# ══════════════ 5 语义缓存 ══════════════
s = slide(); eyebrow(s, "05", "实测结论 · 缓存")
heading(s, "语义缓存的误命中，是把别人问题的答案自信地回给这个用户", size=31)
card(s, ML, 2.3, 5.2, 1.75, "应当命中的最低相似度", None, None)
bignum(s, ML + 0.3, 2.62, "0.651", "同义改写，该复用答案", VERIFY, 40)
card(s, ML + 5.6, 2.3, 5.2, 1.75, "不该命中的最高相似度", None, None, CAVEAT)
bignum(s, ML + 5.9, 2.62, "0.954", "「合同盖公章找谁」vs「合同盖合同章找谁」", CAVEAT, 40)
body(s, "两组完全重叠，所以不存在任何一个阈值能把它们分开。阈值调到 0.94 时缓存已经基本不工作"
        "（10 条同义改写只认出 1 条），却仍然有一次误命中。", 4.35, size=15)
body(s, "这不是模型不好，是「语义相似」和「能用同一个答案」本来就是两回事。"
        "书里给的解法是三层：稠密相似度找回 → 标识符护栏分辨 → 模型裁判兜底。", 5.25, size=15)
quote(s, 6.2, "顺带一条：让 0.5B 小模型当等价性裁判，它对六条测试全答「是」——"
               "不是判错，是根本没在判。这种恒定输出式的失效比判错更难发现。")

# ══════════════ 6 中文 BM25 ══════════════
s = slide(); eyebrow(s, "06", "实测结论 · 中文")
heading(s, "「中文 BM25 必须用 jieba 分词」——一半对，一半错")
body(s, "这是网上流传很广、却几乎找不到量化验证的说法。书里在 1241 篇中文语料上"
        "只换分词方式、其余全部固定，量了一遍：", 2.3, size=15)
table(s, ML, 3.15, [
    ["分词方式", "Recall@5", "结论"],
    ["不分词（按空格切）", "10 / 59", "「不分词召回骤降」是真的，比想象的更极端"],
    ["单字 unigram", "47 / 59", "—"],
    ["字符 bigram", "57 / 59", "反而是最好的"],
    ["jieba 分词", "54 / 59", "不如 bigram"],
], [3.0, 2.0, 5.6], hl=3)
body(s, "正确的说法不是「必须用 jieba」，而是「必须切得足够细，具体用哪种切法要在你自己的语料上量」。"
        "把方向当成做法，是这类「听起来有道理」的说法最常见的害人方式。", 5.75, size=14.5)

# ══════════════ 7 离线交付 ══════════════
s = slide(); eyebrow(s, "07", "实测结论 · 交付")
heading(s, "「这个工具支持离线」要拆成两问分别验")
body(s, "物理隔离交付里最反直觉的一条：用 cosign 给镜像签名，签能离线，验默认不能。", 2.3, size=16)
card(s, ML, 3.0, 10.9, 1.25, "断网环境实测", None, None, CAVEAT)
tf = tb(s, ML + 0.3, 3.42, 10.3, 0.8)
line(tf, 'Error: getting trusted root from TUF …', 12.5, CAVEAT, mono=True, first=True)
line(tf, 'Get "https://tuf-repo-cdn.sigstore.dev/15.root.json": network is unreachable',
     12.5, CAVEAT, mono=True, space_before=3)
body(s, "明明用的是自己生成的密钥对，验签却仍然要去联网拉信任根——而且失败发生在验签之前，"
        "连「签名对不对」都还没开始算。解法是随包多带一份本地信任根（实测只有 73 字节）。", 4.55, size=15)
quote(s, 5.6, "交付现场断网的是客户那一侧，也就是验签那一侧——恰恰是更容易出问题的那一半。"
               "真到现场才发现验不了签，补救的余地非常小。")

# ══════════════ 8 vLLM 显存 ══════════════
s = slide(); eyebrow(s, "08", "实测结论 · 推理")
heading(s, "vLLM 的 gpu_memory_utilization 是预留比例，不是按需上限", size=30)
bignum(s, ML, 2.3, "21.2 GB", None, VERIFY, 48)
tf = tb(s, ML + 3.4, 2.5, 7.5, 1.0)
line(tf, "在 24.5GB 的 RTX 4090 上，加载一个 0.5B 小模型时的实际占用。",
     15, INK2, first=True, spacing=1.4)
body(s, "vLLM 按比例把显存整块划走，权重用不完的部分全部转成 KV Cache。"
        "这不是浪费，是设计——KV Cache 越大能并发的请求越多。但它意味着两件事：", 3.6, size=15)
for i, t in enumerate(["你没法靠换小模型来省显存。",
                       "同一张卡上起第二个实例会直接失败，因为第一个已经把 85% 划走了。"]):
    tf = tb(s, ML + 0.3, 4.55 + i * 0.45, 10.2, 0.4)
    line(tf, "·  " + t, 15, INK, first=True)
body(s, "同一次实测还量到：模型加载耗时 130.3 秒。这个数直接决定 K8s 探针怎么配——"
        "沿用 Web 服务惯用的几秒超时，会在模型还没加载完时就把 Pod 判死并反复重启。", 5.7, size=14.5)

# ══════════════ 9 工具描述 ══════════════
s = slide(); eyebrow(s, "09", "安全边界")
heading(s, "工具描述写得好不好，直接决定 Agent 会不会误触发不可逆操作", size=31)
body(s, "模型没读过你的代码，它判断「该不该批这单」的全部依据，就是工具的那几段文字。", 2.3, size=16)
card(s, ML, 3.0, 5.2, 1.85, "工具选择准确率", None, None)
bignum(s, ML + 0.3, 3.35, "10/10", "连测三轮稳定", VERIFY, 40)
card(s, ML + 5.6, 3.0, 5.2, 1.85, "危险工具误触发", None, None)
bignum(s, ML + 5.9, 3.35, "0", "三条诱导性表述一次都没触发审批", VERIFY, 40)
body(s, "起作用的是描述里那三条带反面例句的排除项。只写「仅在明确授权时调用」是不够的——"
        "模型对「什么算明确」的理解和你不一定一致，例句才把边界钉死。", 5.2, size=15)
body(s, "这个数字要谨慎解读：它说明在这套描述下、对这个模型、这 10 条用例成立，"
        "不是「这样写就绝对安全」的保证。该记住的是方法——把误触发做成可重复运行的检查。",
     6.15, size=13.5, color=MUTED)

# ══════════════ 10 案例 ══════════════
s = slide(); eyebrow(s, "10", "案例")
heading(s, "七个案例，其中一个专讲失败")
cases = [("案例 A", "物理隔离主权 AI", "断网环境的完整交付链路"),
         ("案例 B", "MCP 盘活遗留系统", "只认 mTLS 的老 ERP 怎么接进 Agent"),
         ("案例 C", "混乱数据与自主对账", "脏票据摄取 + 后台对账智能体"),
         ("案例 D", "工业质检", "边缘推理与 MES/SCADA 对接"),
         ("案例 E", "智慧矿山", "弱网下的本地告警与延迟同步"),
         ("案例 F", "高并发 RAG 性能手术", "延迟预算怎么拆、怎么守"),
         ("案例 G", "失败尸检集", "五个真实死法逐个复盘")]
for i, (k, t, d) in enumerate(cases):
    r, c = divmod(i, 4)
    card(s, ML + c * 2.85, 2.35 + r * 1.85, 2.6, 1.6, k, t, d,
         CAVEAT if k == "案例 G" else None)
tf = tb(s, ML + 2.85 * 3, 4.2 + 1.85, 2.6, 1.0)
line(tf, "项目是怎么死的，比它怎么活的更有信息量。", 13, INK2, first=True, spacing=1.4)

# ══════════════ 11 中国落地 ══════════════
s = slide(); eyebrow(s, "11", "中国落地")
heading(s, "国产栈选型的第一关不是技术，是法务")
body(s, "书里把这条列为该章「最需要记住的一条硬约束」：Qwen2.5-3B 用的是研究许可、不可商用；"
        "而同系列的 72B 用的是允许商用的许可。档位相邻不代表授权条款相邻。", 2.3, size=15.5)
body(s, "这个坑和所有技术直觉都反着来：3B 参数最小、显存最省，在资源受限的现场看起来最合适"
        "——恰恰是它不能商用。选错的代价不是效果差一点，是项目上线后被叫停。", 3.35, size=15.5)
card(s, ML, 4.5, 5.2, 1.65, "五步判据", None,
     "商用授权 → 显存预算 → 上下文需求 → 函数调用 → 中文能力。顺序不能换。")
card(s, ML + 5.6, 4.5, 5.2, 1.65, "走完为空怎么办", None,
     "可行集为空说明约束本身有冲突，应回去和客户谈判放松某个维度，不是硬凑方案上线。")
body(s, "这一部还包含信创改造、等保与备案实操、国企 / 银行 / 政务的交付差异——"
        "都是海外资料里查不到的内容。", 6.45, size=13.5, color=MUTED)

# ══════════════ 12 Lab ══════════════
s = slide(); eyebrow(s, "12", "动手")
heading(s, "12 个 Lab，每个都有一份验收脚本")
body(s, "verify.sh 全项通过且退出码为 0 才算完成。这不只是给读者的练习——"
        "正文里的代码必须先在 Lab 里跑通，才允许标注「已实测」。", 2.3, size=15.5)
labs = [("Lab-01 / 04 / 06", "脏 PDF 摄取管道、混合检索与质量诊断、私有化推理与语义缓存"),
        ("Lab-02 / 03 / 10", "mTLS + OAuth2 桥接、遗留 API 封装为 MCP、事件流对账智能体"),
        ("Lab-05 / 08", "企业 Evals 框架与 CI 门禁、可观测性与成本看板"),
        ("Lab-07 / 09", "自包含离线部署包、国产栈演练"),
        ("Lab-11 / 12", "工作流发现演练、毕业项目（唯一一次独立走完六道关卡）")]
for i, (k, t) in enumerate(labs):
    r, c = divmod(i, 3)
    card(s, ML + c * 3.75, 3.3 + r * 1.55, 3.5, 1.35, k, None, t)
card(s, ML + 3.75 * 2, 3.3 + 1.55, 3.5, 1.35, "这件事的副作用", None,
     "做 Lab 反复暴露出正文与实现的矛盾——这类错误光靠通读审不出来。", VERIFY)
quote(s, 6.45, "举一个：某章正文规定「高危阻断、中低危降级」，而同章代码在做整体阻断。"
                "再举一个：某段 schema 强制「金额必须为正」，而红冲发票的金额本来就是负的。")

# ══════════════ 13 事实核查 ══════════════
s = slide(); eyebrow(s, "13", "事实核查")
heading(s, "这本书是 AI 参与写的，所以核查做得比一般书更狠")
for i, (v, k, t, col) in enumerate([
        ("530", "逐条核对的引文", "全量核对，不抽样", VERIFY),
        ("87", "被修正 / 删除", "失败率 16.4%", CAVEAT),
        ("27→3", "未实测代码标注", "做完 Lab 后从 27 处降到 3 处", VERIFY)]):
    card(s, ML + i * 3.75, 2.3, 3.5, 1.75, k, None, None)
    bignum(s, ML + i * 3.75 + 0.28, 2.62, v, t, col, 38, w=3.0)
body(s, "最初只抽样审了一个文件，查出 3 处伪造引文；同一个文件做全量核对，查出 27 处"
        "——抽样漏检率 89%。这件事之后，全书改为全量核对，并把它设成写作前的硬门槛。", 4.35, size=15)
body(s, "被删掉的包括：不存在的合同条款、编造的法规原文、把甲公司的客户写成乙公司的案例、"
        "一个查无来源却被反复引用的行业数字。", 5.3, size=15)
body(s, "剩下那 3 处不是遗漏——它们超出任何 Lab 的验证范围，附录里逐处写明缺的是什么条件。"
        "「你有 GPU 就能验」和「不知道能不能用」对读者是两回事。", 6.25, size=13.5, color=MUTED)

# ══════════════ 14 结语 ══════════════
s = slide(PANEL); eyebrow(s, None, "结语")
heading(s, "这本书真正想教的，是一种判断力", y=1.5, size=40)
body(s, "不是记住多少个数字，而是养成一个习惯：看到一个结论，"
        "先问它是怎么量出来的，以及在什么条件下量的。", 3.0, size=18)
body(s, "书里几乎每一条实测结论，最初都有一个「听起来很有道理」的错误版本"
        "——包括作者自己写下的。改对它们的方式只有一个：把代码跑起来。", 4.1, size=16)
body(s, "方法论那一半用的是同一条规矩：22 个失败案例全部追到一手文书——"
        "审计报告、监管裁定、议会调查、同行评审论文，每一条都标了核到哪一层、哪些数字不能引。", 4.85, size=14.5, color=MUTED)
ln = s.shapes.add_shape(1, In(ML), In(5.3), In(11.2), In(0.012))
ln.fill.solid(); ln.fill.fore_color.rgb = RULE
ln.line.fill.background(); ln.shadow.inherit = False
for i, (v, k) in enumerate([("6", "道关卡"), ("22", "个失败案例"),
                            ("12", "个 Lab"), ("349", "条参考来源")]):
    tf = tb(s, ML + i * 2.5, 5.55, 2.3, 0.9)
    line(tf, v, 28, INK, bold=True, mono=True, first=True)
    line(tf, k, 11, MUTED, space_before=2)
tf = tb(s, ML, 6.75, 11, 0.4)
line(tf, "whiteyoungy · 2026 · 本页数字均取自书中实测记录与案例库台账",
     11.5, MUTED, first=True)

import pathlib
out = str(pathlib.Path(__file__).resolve().parent.parent / "FDE-把AI交付到真实世界-内容导览.pptx")
prs.save(out)
print(f"已生成：{out}")
print(f"共 {len(prs.slides.__iter__.__self__._sldIdLst)} 页，16:9")
